"""Build the integrated Excel workbook for the Western-Europe + mechanism study."""
from __future__ import annotations

import io
import json
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[2]
PR = ROOT / "results/prediction_results"
CR = ROOT / "results/combination_recommendations"
OUT = ROOT / "docs/西欧队列与机制分析_完整数据_20260625.xlsx"

J = lambda p: json.load(io.open(p, encoding="utf-8"))
S = J(PR / "we_healthy/analysis_summary.json")
V = J(PR / "bottomup_validation/validation_summary.json")
BIO = J(PR / "bottomup_validation/bio_only_features.json")
LF = J(PR / "lean_factors/summary.json")
MC = J(PR / "lean_factors/model_comparison.json")
BA = J(PR / "bile_acid/assessment.json")
V2R = J(CR / "synergistic_v2_robust_summary_20260625.json")
V2X = J(CR / "synergistic_v2_relaxed_summary_20260625.json")

CN = {"lean_marker_affinity": "与已确认瘦人菌共现亲和度", "diversity_association": "与群落多样性的关联",
      "spore_former": "芽孢形成能力", "oral_origin": "口腔来源菌", "strict_anaerobe": "严格厌氧",
      "fiber_degrader": "纤维降解能力", "prevalence": "流行率", "cooccurrence_degree": "共现网络中心度",
      "mean_log_abundance": "平均丰度水平"}
DIR = {"lean_enriched": "瘦人富集", "obese_enriched": "肥胖富集"}


def main():
    sh: dict[str, pd.DataFrame] = {}

    sh["00_说明"] = pd.DataFrame({
        "工作表": ["01_研究摘要", "02_方法参数", "03_队列筛选流程", "04_合格队列来源", "05_匹配后样本",
                "06_差异分析_全部物种", "07_稳健差异菌", "08_判别验证", "09_程序一致性_逐物种",
                "10_程序一致性_汇总", "11_机制特征判别力", "12_特征跨属泛化", "13_特征矩阵",
                "14_多样性关联_全部物种", "15_胆汁酸评估", "16_BSH阳性物种",
                "17_组合权重v1v2", "18_v2候选池", "19_v2组合评分", "20_敏感性分析"],
        "内容": ["核心指标一览", "全部统计方法与参数设定", "8304→2049 逐级筛选", "7 项研究人口学",
               "904 例匹配样本明细", "199 物种完整统计结果", "37 个稳健差异菌", "留一研究 AUC 与对照",
               "自下而上程序 vs 实测逐物种", "一致率/κ/精确率", "各机制特征判别力（标注循环项）",
               "各特征组合跨属泛化 AUC", "41 显著物种 × 9 因素取值", "199 物种多样性关联与丰度",
               "胆汁酸各轴评估结果", "BSH 阳性物种归属", "v1 与 v2 权重对照",
               "v2 候选池 16 株", "v2 组合评分结果", "放宽纳入标准的敏感性对比"]})

    sh["01_研究摘要"] = pd.DataFrame({
        "指标": ["合格研究数", "覆盖国家", "瘦人总数", "肥胖总数", "匹配后样本量", "匹配后年龄(瘦/胖)",
               "匹配后女性比(瘦/胖)", "受检物种数", "FDR<0.05 物种数", "稳健差异菌数",
               "菌群 AI AUC", "仅年龄性别 AI AUC",
               "—— 程序验证 ——", "稳健菌一致率", "Cohen's κ", "预测保护精确率", "预测促肥胖精确率",
               "—— 机制特征 ——", "多样性关联 AUC", "仅多样性关联 跨属AUC", "多样性+丙酸 跨属AUC",
               "仅丙酸 跨属AUC", "丙酸潜能→瘦人富集率", "丁酸潜能→瘦人富集率",
               "—— 胆汁酸 ——", "BSH 单轴 AUC", "BSH 跨属 AUC", "多样性+BSH 跨属AUC",
               "—— 组合定标 ——", "v2 候选池(严格)", "v2 候选池(放宽)", "v2 首选组合", "敏感性分析首选是否变化"],
        "数值": [S["eligible_studies"], "/".join(S["countries"]), S["n_lean_total"], S["n_obese_total"],
               S["matched"]["n"], f'{S["matched"]["age_lean"]} / {S["matched"]["age_obese"]}',
               f'{S["matched"]["female_lean"]} / {S["matched"]["female_obese"]}',
               S["n_species_tested"], S["n_fdr_significant"], S["n_robust"],
               S["ai_matched_loso_auc"], S["ai_demographics_only_auc"],
               "", f'{V["concordance"]["稳健差异菌"]["一致率"]:.1%}', V["concordance"]["稳健差异菌"]["kappa"],
               "96% (24/25)", "50% (6/12)",
               "", [r["auc"] for r in LF["single_features"] if r["feature"] == "diversity_association"][0],
               MC["仅多样性关联"], MC["多样性关联+丙酸"], MC["仅丙酸"],
               f'{BIO["propionate_pos_lean_rate"]:.0%} (10/10)', f'{BIO["butyrate_pos_lean_rate"]:.0%} (基线83%)',
               "", [r["auc"] for r in BA["single_axis"] if r["axis"] == "BSH 潜能(0/1/2)"][0],
               BA["cross_genus_auc"]["仅 BSH"], BA["cross_genus_auc"]["多样性+BSH"],
               "", V2R["pool_size"], V2X["pool_size"],
               " + ".join(V2R["top1"]["members"].split("; ")), "否（完全不变）"]})

    sh["02_方法参数"] = pd.DataFrame({
        "方法环节": ["数据类型", "物种定量", "人群", "关键排除", "纳入标准", "研究内匹配",
                 "匹配容差", "匹配后平衡", "物种流行率阈值", "丰度变换", "差异主模型", "多重检验校正",
                 "效应量", "小样本校正", "合并模型", "稳健判定", "机器学习模型", "主验证",
                 "混杂对照", "特征跨类群验证", "多样性指标", "多样性关联算法", "组合打分",
                 "敏感性分析", "统计软件", "随机种子"],
        "设定": ["公开鸟枪法全宏基因组（非16S）", "MetaPhlAn 标记基因法，物种相对丰度(%)",
               "西欧成年人（NLD/GBR/DNK/FRA/DEU/ITA/ESP/SWE/AUT/IRL/LUX）",
               "2型糖尿病(T2D)、糖耐量异常(IGT)、高血压",
               "无疾病；明确瘦/肥胖(排除超重)；年龄性别完整；研究内两组各≥10",
               "同一研究内、同性别、最近邻无放回", "年龄差 ≤5 岁",
               f'年龄 {S["matched"]["age_lean"]} vs {S["matched"]["age_obese"]} 岁；女性比 {S["matched"]["female_lean"]} vs {S["matched"]["female_obese"]}',
               "≥10%", "log10(相对丰度 + 0.001)，伪计数处理零值",
               "OLS: log10丰度 ~ 肥胖 + 年龄 + 性别 + 研究(固定效应)",
               "Benjamini-Hochberg FDR", "Hedges g（标准化均数差）",
               "J = 1 − 3/[4(n1+n2) − 9]", "DerSimonian-Laird 随机效应，报告 95%CI 与 I²",
               "FDR<0.05 且 meta 95%CI 排除0 且两法方向一致",
               "平衡随机森林（500树，深度6，最小叶3，类别权重平衡）",
               "留一研究交叉验证（杜绝批次泄漏）",
               "仅年龄+性别同构模型（应≈0.5）",
               "留一属交叉验证（逻辑回归 L2, C=0.5，标准化）",
               "Shannon H = −Σpi·ln(pi)；另计丰富度/Simpson/Pielou",
               "组内（瘦、胖各自）Spearman(log10丰度, Shannon) 后取均值——避免循环论证",
               "3–5 株穷举；v2权重：多样性0.30/效应0.24/互补0.16/丙酸0.12/可培养0.10/属多样性0.08；BSH归零；芽孢−0.05、口腔−0.10",
               "纳入标准放宽为仅 FDR<0.05（不强制 meta CI 排除0）",
               "Python 3.13 / pandas 2.3 / numpy 2.3 / scikit-learn / statsmodels 0.14 / scipy", "0"]})

    sh["03_队列筛选流程"] = pd.DataFrame(S["cohort_filter_steps"]).rename(columns={"step": "筛选步骤", "n": "剩余样本量"})
    c = pd.read_csv(PR / "we_healthy/cohort_sources.csv")
    c.columns = ["研究队列", "国家", "瘦(n)", "肥胖(n)", "瘦_平均年龄", "肥胖_平均年龄",
                 "瘦_女性比", "肥胖_女性比", "瘦_BMI", "肥胖_BMI", "测序类型"]
    c.loc[len(c)] = ["合计", "—", S["n_lean_total"], S["n_obese_total"], "—", "—", "—", "—", "—", "—", "—"]
    sh["04_合格队列来源"] = c
    m = pd.read_csv(PR / "we_healthy/matched_cohort.csv")
    m.columns = ["样本ID", "研究", "国家", "分组(1=肥胖)", "年龄", "女性(1=是)", "BMI"]
    sh["05_匹配后样本"] = m

    da = pd.read_csv(PR / "we_healthy/differential_abundance_adjusted.csv")
    d2 = da.copy(); d2["瘦人效应(log10)"] = -d2.coef_obese_log10
    d2 = d2[["species", "direction", "瘦人效应(log10)", "fold_change_obese_vs_lean", "p", "fdr",
             "meta_g", "ci_low", "ci_high", "I2_percent", "k_studies", "ci_excludes_zero", "robust",
             "coef_age", "coef_female"]]
    d2.columns = ["物种", "方向", "瘦人效应(log10丰度差)", "倍数变化(胖/瘦)", "P值", "FDR",
                  "meta合并g", "95%CI下限", "95%CI上限", "I²(%)", "研究数", "CI排除0", "稳健",
                  "年龄系数", "性别系数"]
    d2["方向"] = d2["方向"].map(DIR)
    sh["06_差异分析_全部物种"] = d2.sort_values("FDR")
    rob = d2[d2["稳健"] == True]
    sh["07_稳健差异菌"] = pd.concat([rob[rob.方向 == "瘦人富集"].sort_values("meta合并g", ascending=False),
                                rob[rob.方向 == "肥胖富集"].sort_values("meta合并g")])
    sh["08_判别验证"] = pd.DataFrame({
        "模型/留出研究": ["【总体】菌群(199物种)", "【对照】仅年龄+性别"] + list(S["ai_per_study_auc"]),
        "留一研究AUC": [S["ai_matched_loso_auc"], S["ai_demographics_only_auc"]] + [S["ai_per_study_auc"][k] for k in S["ai_per_study_auc"]],
        "说明": ["跨研究可泛化的独立信号", "≈随机，证明混杂已消除"] + ["单研究留出"] * len(S["ai_per_study_auc"])})

    conc = pd.read_csv(PR / "bottomup_validation/concordance_per_species.csv")
    c9 = conc[["species", "observed", "bottomup_pred", "match", "fold_change_obese_vs_lean",
               "fdr", "meta_g", "robust", "mean_abs_shap"]].copy()
    c9.columns = ["物种", "实测方向", "程序预测", "是否一致", "倍数变化(胖/瘦)", "FDR", "meta合并g", "稳健", "SHAP重要性"]
    for col in ["实测方向", "程序预测"]:
        c9[col] = c9[col].map(DIR)
    c9["是否一致"] = c9["是否一致"].map({True: "一致", False: "不一致"})
    sh["09_程序一致性_逐物种"] = c9.sort_values(["稳健", "FDR"], ascending=[False, True])
    rows = [{"比对范围": t, "物种数": V["concordance"][t]["n"],
             "一致率": f'{V["concordance"][t]["一致率"]:.1%}', "Cohen's κ": V["concordance"][t]["kappa"]}
            for t in ["全部受检物种", "FDR<0.05 显著", "稳健差异菌"]]
    c10 = pd.DataFrame(rows)
    c10.loc[len(c10)] = {"比对范围": "预测'保护'时精确率", "物种数": "24/25", "一致率": "96%", "Cohen's κ": "可信"}
    c10.loc[len(c10)] = {"比对范围": "预测'促肥胖'时精确率", "物种数": "6/12", "一致率": "50%", "Cohen's κ": "不可信"}
    sh["10_程序一致性_汇总"] = c10

    f1 = pd.read_csv(PR / "bottomup_validation/feature_discriminative_power.csv")
    f1.columns = ["特征", "单特征AUC", "方向", "瘦人富集菌均值", "肥胖富集菌均值"]
    f1["特征类型"] = f1["特征"].map(lambda x: "源自模型自身(循环)" if x in ("abundance_shap_corr", "mean_abs_shap") else "功能潜能(属级)")
    f2 = pd.DataFrame(LF["single_features"])
    f2["特征"] = f2.feature.map(CN)
    f2 = f2[["特征", "auc", "direction", "lean_mean", "obese_mean", "p_mannwhitney", "fdr"]]
    f2.columns = ["特征", "单特征AUC", "方向", "瘦人富集菌均值", "肥胖富集菌均值", "P值(Mann-Whitney)", "FDR"]
    f2["特征类型"] = f2["特征"].map(lambda x: "部分定义性循环" if "共现亲和度" in x else
                              ("生态学特征" if x in ("与群落多样性的关联", "流行率", "共现网络中心度", "平均丰度水平") else "知识型特征"))
    sh["11_机制特征判别力"] = pd.concat([f1, f2], ignore_index=True)
    g = pd.DataFrame({"特征组合": list(MC), "留一属CV_AUC": [MC[k] for k in MC]})
    g["判定"] = g.留一属CV_AUC.map(lambda x: "可跨属泛化" if x >= 0.8 else ("部分泛化" if x >= 0.6 else "不能泛化(≤随机)"))
    sh["12_特征跨属泛化"] = g.sort_values("留一属CV_AUC", ascending=False)
    fm = pd.read_csv(PR / "lean_factors/lean_factor_matrix.csv")
    fm = fm.rename(columns={"species": "物种", "genus": "属", "direction": "方向", "y": "瘦人富集(1)", **CN})
    fm["方向"] = fm["方向"].map(DIR)
    sh["13_特征矩阵"] = fm

    dv = pd.read_csv(PR / "lean_factors/diversity_association_all_species.csv")
    dv = dv[["species", "diversity_rho", "abund_high_div", "abund_low_div", "fold_high_vs_low",
             "prev_high_div", "prev_low_div", "direction", "fdr", "robust"]]
    dv.columns = ["物种", "多样性关联ρ", "高多样性人群丰度%", "低多样性人群丰度%", "倍数(高/低)",
                  "高多样性流行率", "低多样性流行率", "胖瘦方向", "FDR", "稳健"]
    dv["胖瘦方向"] = dv["胖瘦方向"].map(DIR)
    sh["14_多样性关联_全部物种"] = dv.sort_values("多样性关联ρ", ascending=False)

    ba = pd.DataFrame(BA["single_axis"])
    ba.columns = ["评估轴", "AUC", "方向", "瘦人富集菌均值", "肥胖富集菌均值", "P值"]
    sh["15_胆汁酸评估"] = ba
    bp = pd.read_csv(PR / "bile_acid/bsh_positive_species.csv")
    bp["direction"] = bp["direction"].map(DIR)
    bp = bp.rename(columns={"species": "物种", "genus": "属", "bsh": "BSH潜能", "direction": "方向",
                            "fold_change_obese_vs_lean": "倍数(胖/瘦)", "meta_g": "meta合并g",
                            "fdr": "FDR", "diversity_rho": "多样性关联ρ"})
    sh["16_BSH阳性物种"] = bp

    sh["17_组合权重v1v2"] = pd.DataFrame({
        "评分项": ["与多样性关联", "实测胖瘦效应", "机制互补", "丙酸潜能", "菌株性质(可培养)",
                "属多样性", "BSH降胆固醇", "丁酸(正向项)", "惩罚:芽孢形成", "惩罚:口腔来源"],
        "版本1": [0, 0.10, 0.28, "隐含", 0.14, 0.06, 0.16, "隐含", "—", "—"],
        "版本2": [0.30, 0.24, 0.16, 0.12, 0.10, 0.08, 0.00, "不计", -0.05, -0.10],
        "调整依据": ["AUC 0.920，唯一可跨属泛化(0.878)", "本研究 meta 合并效应量", "仍为设计原则，不再主导",
                 "AUC 0.647，但跨属仅0.321", "基本不变", "组合层面",
                 "AUC 0.517(P=0.85)；跨属0.027；纳入后0.878→0.798", "AUC 0.540 且方向相反",
                 "肥胖富集倾向(71% vs 32%)", "肥胖富集倾向(14% vs 0%)"]})
    pool = pd.read_csv(CR / "synergistic_pool_v2_robust_20260625.csv")
    pool = pool[["species", "genus", "diversity_rho", "meta_g", "fold_change_obese_vs_lean",
                 "propionate", "butyrate", "mucin", "culturability_class", "spore_former", "oral_origin"]]
    pool.columns = ["物种", "属", "多样性关联ρ", "meta合并g", "倍数(胖/瘦)", "丙酸", "丁酸", "黏液",
                    "可培养性", "芽孢形成", "口腔来源"]
    sh["18_v2候选池"] = pool.sort_values("多样性关联ρ", ascending=False)
    cb = pd.read_csv(CR / "synergistic_lean_combinations_v2_robust_20260625.csv").head(200)
    cb.columns = ["排名", "组合成员", "株数", "机制轴", "覆盖轴数", "多样性ρ均值", "效应g均值",
                  "丙酸占比", "可培养均值", "属多样性", "芽孢占比", "v2综合分"]
    sh["19_v2组合评分"] = cb
    sh["20_敏感性分析"] = pd.DataFrame({
        "项目": ["纳入标准", "候选池株数", "候选组合数", "首选组合", "首选综合分",
               "含Akkermansia最佳排名", "含Akkermansia最佳分", "与首选差距", "结论"],
        "严格版(robust)": ["FDR<0.05 且 meta CI排除0 且方向一致", V2R["pool_size"], V2R["n_combinations"],
                       " + ".join(V2R["top1"]["members"].split("; ")), V2R["top1"]["composite_v2"],
                       "不在池中(未通过三重判定)", "—", "—", "—"],
        "放宽版(relaxed)": ["仅 FDR<0.05", V2X["pool_size"], V2X["n_combinations"],
                        " + ".join(V2X["top1"]["members"].split("; ")), V2X["top1"]["composite_v2"],
                        "#9", 0.7471, 0.0160, "首选组合完全不变，结论稳健"]})

    with pd.ExcelWriter(OUT, engine="openpyxl") as w:
        for name, df in sh.items():
            df.to_excel(w, sheet_name=name, index=False)
        hf = Font(bold=True, color="FFFFFF"); hfill = PatternFill("solid", fgColor="2C6FB0")
        for ws in w.book.worksheets:
            for cc in ws[1]:
                cc.font = hf; cc.fill = hfill
                cc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.freeze_panes = "A2"
            for col in ws.columns:
                ln = max((len(str(x.value)) if x.value is not None else 0) for x in col)
                ws.column_dimensions[col[0].column_letter].width = min(max(ln + 2, 10), 52)
    print("wrote", OUT)
    for name, df in sh.items():
        print(f"  {name}: {len(df)} 行")


if __name__ == "__main__":
    main()
